from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import landscape, letter, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..config import APP_DESCRIPTOR, APP_NAME, EXPORT_DIR, LOGO_PATH
from ..database import Database
from ..date_utils import display_date, display_datetime
from .inventory_service import InventoryService
from .settings_service import SettingsService


BRAND = "#0878B9"
BRAND_DARK = "#075985"
ROW_A = "#F8FAFC"
ROW_B = "#EEF3F8"
GRID = "#B8C5D3"
TEXT = "#243447"
SUCCESS = "#E8F5EC"
INFO = "#E8F1FB"
WARNING = "#FFF3D6"
DANGER = "#FDE7E7"


class _NumberedFooterCanvas(canvas.Canvas):
    """Canvas que agrega pie consistente y "Página X de Y"."""

    def __init__(self, *args, footer_text: str = "", **kwargs):
        super().__init__(*args, **kwargs)
        self.footer_text = footer_text
        self._saved_page_states: list[dict] = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self._draw_footer(page_count)
            super().showPage()
        super().save()

    def _draw_footer(self, page_count: int) -> None:
        width, _height = self._pagesize
        self.saveState()
        self.setStrokeColor(colors.HexColor("#D4DCE6"))
        self.setLineWidth(0.5)
        self.line(18 * mm, 13.5 * mm, width - 18 * mm, 13.5 * mm)
        self.setFillColor(colors.HexColor("#64748B"))
        self.setFont("Helvetica", 7.8)
        self.drawString(18 * mm, 8.5 * mm, APP_NAME)
        self.drawCentredString(width / 2, 8.5 * mm, self.footer_text)
        self.drawRightString(width - 18 * mm, 8.5 * mm, f"Página {self._pageNumber} de {page_count}")
        self.restoreState()


class ReportService:
    def __init__(self, db: Database) -> None:
        self.db = db
        self.inventory = InventoryService(db)
        self.settings = SettingsService(db)
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    def sales_rows(self, start: str | None = None, end: str | None = None):
        clauses = ["v.estado='CONFIRMADA'"]
        params: list[object] = []
        if start:
            clauses.append("date(v.fecha_hora,'localtime') >= date(?)")
            params.append(start)
        if end:
            clauses.append("date(v.fecha_hora,'localtime') <= date(?)")
            params.append(end)
        where = " AND ".join(clauses)
        with self.db.connect() as con:
            return con.execute(
                f"""
                SELECT v.id,v.fecha_hora,p.nombre,d.cantidad,p.unidad_control,
                       d.precio_unitario,d.subtotal,u.nombre AS usuario
                  FROM ventas v
                  JOIN detalle_ventas d ON d.venta_id=v.id
                  JOIN productos p ON p.id=d.producto_id
                  JOIN usuarios u ON u.id=v.usuario_id
                 WHERE {where}
                 ORDER BY v.fecha_hora DESC,v.id DESC,d.id
                """,
                params,
            ).fetchall()

    @staticmethod
    def _timestamp() -> str:
        return datetime.now().strftime("%Y%m%d_%H%M%S")

    @staticmethod
    def _period_text(start: str | None, end: str | None) -> str:
        if start and end:
            return f"{display_date(start)} al {display_date(end)}"
        if start:
            return f"Desde {display_date(start)}"
        if end:
            return f"Hasta {display_date(end)}"
        return "Todos los registros disponibles"

    @staticmethod
    def _semantic_movement(row) -> str:
        raw = str(row["tipo"] or "").upper()
        ref = str(row["referencia_tipo"] or "").upper()
        motive = str(row["motivo"] or "").casefold()
        if raw == "ANULACION" and ref == "VENTA":
            return "ANULACIÓN DE VENTA"
        if raw == "AJUSTE" and ref == "ENTRADA" and "anulación de entrada" in motive:
            return "ANULACIÓN DE ENTRADA"
        if raw == "AJUSTE" and "merma por vencimiento" in motive:
            return "MERMA"
        return raw or "MOVIMIENTO"

    @staticmethod
    def _movement_kind(label: str) -> str:
        label = label.upper()
        if label == "ENTRADA":
            return "success"
        if label == "VENTA":
            return "info"
        if "ANULACIÓN" in label:
            return "danger"
        if label == "MERMA":
            return "danger"
        return "warning"

    @staticmethod
    def _footer_text(generated_by: str, generated_at: datetime) -> str:
        return f"Generado por: {generated_by or 'Sistema'} · {generated_at:%d/%m/%Y %H:%M}"

    @staticmethod
    def _styles():
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(
            name="ReportBrand",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=10,
            textColor=colors.HexColor(BRAND),
            spaceAfter=3,
        ))
        styles.add(ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=23,
            textColor=colors.HexColor("#152536"),
            alignment=TA_LEFT,
            spaceAfter=7,
        ))
        styles.add(ParagraphStyle(
            name="ReportMeta",
            parent=styles["Normal"],
            fontSize=8.7,
            leading=12,
            textColor=colors.HexColor(TEXT),
        ))
        styles.add(ParagraphStyle(
            name="ReportSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12.5,
            leading=15,
            textColor=colors.HexColor("#152536"),
            spaceBefore=7,
            spaceAfter=6,
        ))
        styles.add(ParagraphStyle(
            name="Cell",
            parent=styles["Normal"],
            fontSize=7.7,
            leading=9.5,
            textColor=colors.HexColor(TEXT),
        ))
        styles.add(ParagraphStyle(
            name="CellCenter",
            parent=styles["Cell"],
            alignment=TA_CENTER,
        ))
        styles.add(ParagraphStyle(
            name="CellHeader",
            parent=styles["Cell"],
            fontName="Helvetica-Bold",
            textColor=colors.white,
        ))
        styles.add(ParagraphStyle(
            name="CellHeaderCenter",
            parent=styles["CellHeader"],
            alignment=TA_CENTER,
        ))
        return styles

    @staticmethod
    def _pdf_logo(max_width: float = 38 * mm, max_height: float = 25 * mm):
        if not LOGO_PATH.exists():
            return None
        try:
            reader = ImageReader(str(LOGO_PATH))
            iw, ih = reader.getSize()
            ratio = min(max_width / iw, max_height / ih)
            return RLImage(str(LOGO_PATH), width=iw * ratio, height=ih * ratio, mask="auto")
        except Exception:
            return None

    def _pdf_header(self, title: str, context: str, styles, available_width: float):
        logo = self._pdf_logo()
        text_stack = [
            Paragraph(title, styles["ReportTitle"]),
            Paragraph(context, styles["ReportMeta"]),
        ]
        if logo is None:
            return text_stack
        brand_width = 43 * mm
        header = Table([[logo, text_stack]], colWidths=[brand_width, max(40 * mm, available_width - brand_width)], hAlign="LEFT")
        header.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (0, 0), 7 * mm),
            ("RIGHTPADDING", (1, 0), (1, 0), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        return [header]

    def _summary_table(self, metadata: list[tuple[str, str]], styles, max_cols: int = 4):
        if not metadata:
            return None
        # Evitar celdas vacías innecesarias: los resúmenes cortos se presentan
        # en una sola fila y los de seis indicadores en dos filas de tres.
        if len(metadata) <= 5:
            max_cols = len(metadata)
        elif len(metadata) == 6:
            max_cols = 3
        cells = []
        for label, value in metadata:
            cells.append(Paragraph(f"<b>{label}</b><br/><font color='#475569'>{value}</font>", styles["ReportMeta"]))
        rows = [cells[i:i + max_cols] for i in range(0, len(cells), max_cols)]
        while rows and len(rows[-1]) < max_cols:
            rows[-1].append("")
        table = Table(rows, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F8FC")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D6E0EA")),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E1E8F0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        return table

    def _pdf_section(
        self,
        title: str,
        headers: list[str],
        rows: list[list[object]],
        styles,
        *,
        widths: list[float] | None = None,
        row_kinds: list[str | None] | None = None,
        centered: set[int] | None = None,
    ) -> list:
        elements: list = [Paragraph(title, styles["ReportSection"])]
        centered = centered or set()
        processed: list[list[object]] = []
        for row in rows:
            processed.append([
                Paragraph("" if v is None else str(v), styles["CellCenter"] if idx in centered else styles["Cell"])
                for idx, v in enumerate(row)
            ])
        data = [[Paragraph(str(h), styles["CellHeaderCenter"] if idx in centered else styles["CellHeader"]) for idx, h in enumerate(headers)]] + processed
        table = Table(data, repeatRows=1, colWidths=widths, hAlign="LEFT")
        commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor(ROW_A), colors.HexColor(ROW_B)]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(GRID)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4.5),
            ("TOPPADDING", (0, 0), (-1, -1), 4.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
        ]
        if row_kinds:
            kind_fill = {
                "success": colors.HexColor(SUCCESS),
                "info": colors.HexColor(INFO),
                "warning": colors.HexColor(WARNING),
                "danger": colors.HexColor(DANGER),
            }
            # Resaltar únicamente la celda "Tipo" si existe; si no, toda la fila suavemente.
            type_col = next((i for i, h in enumerate(headers) if h == "Tipo"), None)
            for idx, kind in enumerate(row_kinds, start=1):
                if kind in kind_fill:
                    target = (type_col, idx) if type_col is not None else (0, idx)
                    end = target if type_col is not None else (-1, idx)
                    commands.append(("BACKGROUND", target, end, kind_fill[kind]))
                    if type_col is not None:
                        commands.append(("FONTNAME", target, end, "Helvetica-Bold"))
        table.setStyle(TableStyle(commands))
        elements.extend([table, Spacer(1, 8 * mm)])
        return elements

    def _to_pdf(
        self,
        title: str,
        context: str,
        metadata: list[tuple[str, str]],
        sections: list[dict],
        *,
        generated_by: str,
        filename: str | None = None,
        pagesize=landscape(letter),
    ) -> Path:
        generated_at = datetime.now()
        path = EXPORT_DIR / (filename or f"{title.lower().replace(' ','_')}_{self._timestamp()}.pdf")
        doc = SimpleDocTemplate(
            str(path),
            pagesize=pagesize,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=17 * mm,
            bottomMargin=19 * mm,
            title=title,
            author=generated_by or "Sistema",
        )
        styles = self._styles()
        elements = self._pdf_header(title, context, styles, doc.width)
        elements.append(Spacer(1, 4 * mm))
        summary = self._summary_table(metadata, styles)
        if summary is not None:
            elements.extend([summary, Spacer(1, 5 * mm)])
        for section in sections:
            elements.extend(self._pdf_section(
                section["title"], section["headers"], section["rows"], styles,
                widths=section.get("widths"), row_kinds=section.get("row_kinds"), centered=section.get("centered"),
            ))
        footer = self._footer_text(generated_by, generated_at)
        doc.build(
            elements,
            canvasmaker=lambda *args, **kwargs: _NumberedFooterCanvas(*args, footer_text=footer, **kwargs),
        )
        return path

    def _excel_sheet(
        self,
        ws,
        title: str,
        context: str,
        metadata: list[tuple[str, str]],
        headers: list[str],
        rows: list[list[object]],
        *,
        generated_by: str,
        generated_at: datetime,
    ) -> int:
        max_col = max(1, len(headers))
        title_col = 3 if max_col >= 4 else 2
        if LOGO_PATH.exists():
            try:
                logo = XLImage(str(LOGO_PATH))
                ratio = min(118 / max(1, logo.width), 82 / max(1, logo.height))
                logo.width = int(logo.width * ratio)
                logo.height = int(logo.height * ratio)
                ws.add_image(logo, "A1")
            except Exception:
                ws.cell(1, 1, APP_NAME).font = Font(size=12, bold=True, color=BRAND.lstrip("#"))
        else:
            ws.cell(1, 1, APP_NAME).font = Font(size=12, bold=True, color=BRAND.lstrip("#"))

        ws.row_dimensions[1].height = 42
        ws.row_dimensions[2].height = 24
        ws.merge_cells(start_row=1, start_column=title_col, end_row=1, end_column=max_col)
        ws.cell(1, title_col, title)
        ws.cell(1, title_col).font = Font(size=18, bold=True, color="152536")
        ws.cell(1, title_col).alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=2, start_column=title_col, end_row=2, end_column=max_col)
        ws.cell(2, title_col, context)
        ws.cell(2, title_col).font = Font(size=10, color="475569")
        ws.cell(2, title_col).alignment = Alignment(vertical="top", wrap_text=True)
        if title_col > 1:
            ws.merge_cells(start_row=3, start_column=title_col, end_row=3, end_column=max_col)
            ws.cell(3, title_col, APP_DESCRIPTOR)
            ws.cell(3, title_col).font = Font(size=9, color=BRAND.lstrip("#"), italic=True)

        row = 5
        for label, value in metadata:
            ws.cell(row, 1, label).font = Font(bold=True, color="334155")
            ws.cell(row, 2, value)
            row += 1
        row += 1
        header_row = row
        for col, text in enumerate(headers, 1):
            cell = ws.cell(header_row, col, text)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=BRAND.lstrip("#"))
            cell.alignment = Alignment(horizontal="center", vertical="center")
        thin = Side(style="thin", color="C7D2DF")
        for values in rows:
            row += 1
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, value)
                cell.fill = PatternFill("solid", fgColor=(ROW_A if row % 2 == 0 else ROW_B).lstrip("#"))
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                header = headers[col - 1]
                if header in {"Precio", "Precio unitario", "Subtotal", "Total"} and isinstance(value, (int, float)):
                    cell.number_format = '$#,##0.00'
                if header in {"Cantidad", "Stock total", "Stock vendible", "Bloqueado", "Vencido", "Stock mínimo", "Días restantes"}:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max(row, header_row)}"
        ws.freeze_panes = f"A{header_row + 1}"
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = APP_NAME
        ws.oddFooter.center.text = f"Generado por: {generated_by or 'Sistema'} · {generated_at:%d/%m/%Y %H:%M}"
        ws.oddFooter.right.text = "Página &P de &N"
        ws.oddFooter.left.size = 8
        ws.oddFooter.center.size = 8
        ws.oddFooter.right.size = 8
        ws.page_margins.left = 0.45
        ws.page_margins.right = 0.45
        ws.page_margins.top = 0.55
        ws.page_margins.bottom = 0.65
        for col_idx in range(1, max_col + 1):
            values = [ws.cell(r, col_idx).value for r in range(1, min(row, 250) + 1)]
            max_len = max((len(str(v)) if v is not None else 0) for v in values)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 11), 42)
        return row

    def _to_excel(
        self,
        title: str,
        context: str,
        metadata: list[tuple[str, str]],
        headers: list[str],
        rows: list[list[object]],
        *,
        generated_by: str,
        filename: str | None = None,
    ) -> Path:
        path = EXPORT_DIR / (filename or f"{title.lower().replace(' ','_')}_{self._timestamp()}.xlsx")
        generated_at = datetime.now()
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        self._excel_sheet(ws, title, context, metadata, headers, rows, generated_by=generated_by, generated_at=generated_at)
        wb.save(path)
        return path

    def export_inventory(self, fmt: str = "xlsx", generated_by: str = "") -> Path:
        records = self.inventory.inventory_rows()
        headers = ["Producto", "Unidad", "Precio", "Stock total", "Stock vendible", "Bloqueado", "Vencido", "Stock mínimo", "Ubicación"]
        rows = [[
            r["nombre"], r["unidad_control"], float(r["precio"]), int(r["stock_total"]), int(r["stock_vendible"]),
            int(r["stock_bloqueado"]), int(r["stock_vencido"]), int(r["stock_minimo"]), r["ubicacion"] or "",
        ] for r in records]
        low_count = sum(1 for r in records if int(r["stock_vendible"]) <= int(r["stock_minimo"]))
        expired_lots = len(self.inventory.expired())
        blocked_units = sum(int(r["stock_bloqueado"]) for r in records)
        metadata = [
            ("Fecha de corte", datetime.now().strftime("%d/%m/%Y %H:%M")),
            ("Productos", str(len(rows))),
            ("Stock bajo", str(low_count)),
            ("Unidades bloqueadas", str(blocked_units)),
            ("Lotes vencidos", str(expired_lots)),
        ]
        context = "Estado actual del inventario al momento de generar el documento."
        if fmt == "xlsx":
            return self._to_excel("Inventario actual", context, metadata, headers, rows, generated_by=generated_by)
        widths = [42*mm, 18*mm, 18*mm, 19*mm, 22*mm, 20*mm, 18*mm, 21*mm, 30*mm]
        return self._to_pdf(
            "Inventario actual", context, metadata,
            [{"title": "Detalle de inventario", "headers": headers, "rows": [[r[0], r[1], f"${r[2]:.2f}", *r[3:]] for r in rows], "widths": widths, "centered": {1,2,3,4,5,6,7}}],
            generated_by=generated_by,
        )

    def export_sales(self, start: str | None, end: str | None, fmt: str = "xlsx", generated_by: str = "") -> Path:
        records = self.sales_rows(start, end)
        headers = ["Venta", "Fecha", "Producto", "Cantidad", "Unidad", "Precio unitario", "Subtotal", "Usuario"]
        rows = [[
            int(r["id"]), display_datetime(r["fecha_hora"]), r["nombre"], int(r["cantidad"]), r["unidad_control"],
            float(r["precio_unitario"]), float(r["subtotal"]), r["usuario"],
        ] for r in records]
        total = round(sum(float(r["subtotal"]) for r in records), 2)
        sales_count = len({int(r["id"]) for r in records})
        metadata = [
            ("Período", self._period_text(start, end)),
            ("Ventas confirmadas", str(sales_count)),
            ("Líneas de detalle", str(len(rows))),
            ("Total vendido", f"${total:,.2f}"),
        ]
        context = "Ventas confirmadas dentro del período seleccionado."
        if fmt == "xlsx":
            return self._to_excel("Reporte de ventas", context, metadata, headers, rows, generated_by=generated_by)
        widths = [16*mm, 30*mm, 48*mm, 19*mm, 20*mm, 25*mm, 23*mm, 30*mm]
        return self._to_pdf(
            "Reporte de ventas", context, metadata,
            [{"title": "Detalle de ventas", "headers": headers, "rows": [[*r[:5], f"${r[5]:.2f}", f"${r[6]:.2f}", r[7]] for r in rows], "widths": widths, "centered": {0,1,3,4,5,6}}],
            generated_by=generated_by,
        )

    def export_movements(self, start: str | None, end: str | None, fmt: str = "xlsx", generated_by: str = "") -> Path:
        records = self.inventory.movements(start, end)
        semantic = [self._semantic_movement(r) for r in records]
        counts = Counter(semantic)
        headers = ["Fecha", "Producto", "Unidad", "Lote", "Tipo", "Cantidad", "Usuario", "Motivo"]
        rows = []
        for r, label in zip(records, semantic):
            qty = int(r["cantidad"])
            rows.append([
                display_datetime(r["fecha_hora"]), r["nombre"], r["unidad_control"], r["codigo_interno"] or "-",
                label, f"{qty:+d}", r["usuario"], r["motivo"] or "",
            ])
        summary_parts = []
        for key in ("ENTRADA", "VENTA", "AJUSTE", "MERMA", "ANULACIÓN DE VENTA", "ANULACIÓN DE ENTRADA"):
            if counts.get(key):
                summary_parts.append(f"{key.title()}: {counts[key]}")
        metadata = [
            ("Período", self._period_text(start, end)),
            ("Movimientos", str(len(rows))),
            ("Resumen", " · ".join(summary_parts) or "Sin movimientos"),
        ]
        context = "Trazabilidad cronológica de entradas, ventas, ajustes, anulaciones y mermas."
        if fmt == "xlsx":
            return self._to_excel("Historial de movimientos", context, metadata, headers, rows, generated_by=generated_by)
        widths = [27*mm, 42*mm, 18*mm, 25*mm, 32*mm, 18*mm, 27*mm, 56*mm]
        return self._to_pdf(
            "Historial de movimientos", context, metadata,
            [{
                "title": "Detalle de movimientos", "headers": headers, "rows": rows, "widths": widths,
                "row_kinds": [self._movement_kind(v) for v in semantic], "centered": {0,2,3,4,5,6},
            }],
            generated_by=generated_by,
            pagesize=landscape(letter),
        )

    def export_expirations(self, days: int = 90, fmt: str = "xlsx", generated_by: str = "") -> Path:
        block_days = self.settings.get_block_days()
        all_upcoming = list(self.inventory.expiring(days))
        blocked = [r for r in all_upcoming if int(r["dias_restantes"]) <= block_days]
        upcoming = [r for r in all_upcoming if int(r["dias_restantes"]) > block_days]
        expired = list(self.inventory.expired())

        upcoming_headers = ["Producto", "Unidad", "Lote", "Cantidad", "Vencimiento", "Días restantes"]
        upcoming_rows = [[r["nombre"], r["unidad_control"], r["codigo_interno"], int(r["cantidad_actual"]), display_date(r["fecha_vencimiento"]), int(r["dias_restantes"])] for r in upcoming]
        blocked_headers = ["Producto", "Unidad", "Lote", "Cantidad", "Vencimiento", "Estado"]
        blocked_rows = [[
            r["nombre"], r["unidad_control"], r["codigo_interno"], int(r["cantidad_actual"]), display_date(r["fecha_vencimiento"]),
            "Vence hoy / bloqueado" if int(r["dias_restantes"]) == 0 else f"Bloqueado ({int(r['dias_restantes'])} días)",
        ] for r in blocked]
        expired_headers = ["Producto", "Unidad", "Lote", "Cantidad", "Venció"]
        expired_rows = [[r["nombre"], r["unidad_control"], r["codigo_interno"], int(r["cantidad_actual"]), display_date(r["fecha_vencimiento"])] for r in expired]
        metadata = [
            ("Fecha de corte", datetime.now().strftime("%d/%m/%Y")),
            ("Seguimiento", f"Próximos {days} días"),
            ("Política de venta", f"Bloqueo con {block_days} días o menos"),
            ("Próximos", str(len(upcoming_rows))),
            ("Bloqueados", str(len(blocked_rows))),
            ("Vencidos", str(len(expired_rows))),
        ]
        context = "Control de lotes próximos a vencer, bloqueados para venta y ya vencidos."
        if fmt == "xlsx":
            path = EXPORT_DIR / f"vencimientos_{self._timestamp()}.xlsx"
            generated_at = datetime.now()
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Próximos a vencer"
            self._excel_sheet(ws1, "Próximos a vencer", context, metadata, upcoming_headers, upcoming_rows, generated_by=generated_by, generated_at=generated_at)
            ws2 = wb.create_sheet("Bloqueados")
            self._excel_sheet(ws2, "Lotes bloqueados para venta", context, metadata, blocked_headers, blocked_rows, generated_by=generated_by, generated_at=generated_at)
            ws3 = wb.create_sheet("Vencidos")
            self._excel_sheet(ws3, "Productos vencidos", context, metadata, expired_headers, expired_rows, generated_by=generated_by, generated_at=generated_at)
            wb.save(path)
            return path

        sections = [
            {"title": "Próximos a vencer", "headers": upcoming_headers, "rows": upcoming_rows, "widths": [40*mm,20*mm,27*mm,20*mm,28*mm,25*mm], "centered": {1,2,3,4,5}},
            {"title": "Lotes bloqueados para venta", "headers": blocked_headers, "rows": blocked_rows, "widths": [40*mm,20*mm,27*mm,20*mm,28*mm,40*mm], "row_kinds": ["warning"]*len(blocked_rows), "centered": {1,2,3,4,5}},
            {"title": "Productos vencidos", "headers": expired_headers, "rows": expired_rows, "widths": [48*mm,22*mm,30*mm,22*mm,30*mm], "row_kinds": ["danger"]*len(expired_rows), "centered": {1,2,3,4}},
        ]
        return self._to_pdf("Reporte de vencimientos", context, metadata, sections, generated_by=generated_by, pagesize=portrait(letter))
