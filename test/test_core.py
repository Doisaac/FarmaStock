from __future__ import annotations

import tempfile
import unittest
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.database import Database
from app.services.auth_service import AuthService
from app.services.entry_service import EntryService
from app.services.inventory_service import InventoryService
from app.services.product_service import ProductService
from app.services.sale_service import SaleService
from app.services.settings_service import SettingsService


class CoreFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "test.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.session = self.auth.login("admin", "Admin123!")
        assert self.session is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.sales = SaleService(self.db)
        self.inventory = InventoryService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_product_entry_sale_and_cancel(self):
        pid = self.products.create_product("Palagrip", "Blíster", 4.0, 5, "Vitrina 2")
        first = self.entries.register_entry(pid, 2, "2026-12-15", self.session.id)
        second = self.entries.register_entry(pid, 20, "2027-03-20", self.session.id)
        self.assertEqual(first["internal_code"], f"P{pid:04d}-L001")
        self.assertEqual(second["internal_code"], f"P{pid:04d}-L002")

        sale_id = self.sales.create_sale([{"product_id": pid, "quantity": 3}], self.session.id)
        lots = self.inventory.lots_for_product(pid)
        self.assertEqual([l["cantidad_actual"] for l in lots], [0, 19])

        self.sales.cancel_sale(sale_id, "Prueba de anulación", self.session.id)
        lots = self.inventory.lots_for_product(pid)
        self.assertEqual([l["cantidad_actual"] for l in lots], [2, 20])

    def test_expired_lot_is_not_sold(self):
        pid = self.products.create_product("Medicamento vencido", "Unidad", 1.0, 0)
        self.entries.register_entry(pid, 5, "2020-01-01", self.session.id)
        with self.assertRaises(ValueError):
            self.sales.create_sale([{"product_id": pid, "quantity": 1}], self.session.id)

    def test_adjustment_leaves_trace(self):
        pid = self.products.create_product("Gasas", "Unidad", 0.25, 10)
        entry = self.entries.register_entry(pid, 30, None, self.session.id)
        aid = self.inventory.adjust_lot(entry["lot_id"], 28, "Conteo físico", self.session.id)
        self.assertGreater(aid, 0)
        lot = self.inventory.lots_for_product(pid)[0]
        self.assertEqual(lot["cantidad_actual"], 28)
        movements = self.inventory.movements(product_id=pid)
        self.assertEqual(movements[0]["tipo"], "AJUSTE")
        self.assertEqual(movements[0]["cantidad"], -2)

    def test_low_stock_and_backup_restore(self):
        from app.services.backup_service import BackupService
        pid = self.products.create_product("Rapidol", "Blíster", 3.5, 10)
        self.entries.register_entry(pid, 8, "2027-12-31", self.session.id)
        low = self.inventory.low_stock()
        self.assertTrue(any(r["nombre"] == "Rapidol" for r in low))

        backup_service = BackupService(self.db)
        backup_path = backup_service.create_backup(Path(self.tmp.name) / "backup.db")
        self.products.create_product("Temporal", "Unidad", 1.0, 0)
        self.assertTrue(any(r["nombre"] == "Temporal" for r in self.products.list_products()))
        backup_service.restore_backup(backup_path)
        self.assertFalse(any(r["nombre"] == "Temporal" for r in self.products.list_products()))

class ExtendedFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "extended.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.inventory = InventoryService(self.db)
        self.sales = SaleService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_roles_and_last_admin_protection(self):
        uid = self.auth.create_user("Operador", "operador", "Clave123", "OPERADOR")
        session = self.auth.login("operador", "Clave123")
        self.assertIsNotNone(session)
        self.assertEqual(session.rol, "OPERADOR")
        with self.assertRaises(ValueError):
            self.auth.set_active(self.admin.id, False)
        self.auth.set_active(uid, False)
        self.assertIsNone(self.auth.login("operador", "Clave123"))

    def test_entry_cancel_only_without_later_movements(self):
        pid = self.products.create_product("Entrada anulable", "Unidad", 1.0, 0)
        entry = self.entries.register_entry(pid, 10, "2027-12-31", self.admin.id)
        self.entries.cancel_entry(entry["entry_id"], "Error de digitación", self.admin.id)
        lot = self.inventory.lots_for_product(pid)[0]
        self.assertEqual(lot["cantidad_actual"], 0)
        rows = self.entries.list_entries()
        self.assertEqual(rows[0]["estado"], "ANULADA")

        pid2 = self.products.create_product("Entrada con venta", "Unidad", 2.0, 0)
        entry2 = self.entries.register_entry(pid2, 10, "2027-12-31", self.admin.id)
        self.sales.create_sale([{"product_id": pid2, "quantity": 1}], self.admin.id)
        with self.assertRaises(ValueError):
            self.entries.cancel_entry(entry2["entry_id"], "No debe", self.admin.id)

    def test_discard_expired_creates_adjustment(self):
        pid = self.products.create_product("Vencido merma", "Unidad", 1.0, 0)
        entry = self.entries.register_entry(pid, 5, "2020-01-01", self.admin.id)
        self.inventory.discard_expired(entry["lot_id"], 2, "Retiro físico", self.admin.id)
        lot = self.inventory.lots_for_product(pid)[0]
        self.assertEqual(lot["cantidad_actual"], 3)
        self.assertTrue(self.inventory.movements(product_id=pid)[0]["motivo"].startswith("Merma por vencimiento"))

    def test_report_includes_upcoming_and_expired(self):
        from app.services.report_service import ReportService
        pid = self.products.create_product("Reporte venc", "Unidad", 1.0, 0)
        self.entries.register_entry(pid, 2, "2020-01-01", self.admin.id)
        self.entries.register_entry(pid, 3, "2099-01-01", self.admin.id)
        service = ReportService(self.db)
        path = service.export_expirations(36500, "xlsx", generated_by="Administrador")
        self.assertTrue(path.exists())
        from openpyxl import load_workbook
        wb = load_workbook(path)
        self.assertIn("Próximos a vencer", wb.sheetnames)
        self.assertIn("Vencidos", wb.sheetnames)


class IntegrityRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "regression.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.sales = SaleService(self.db)
        self.inventory = InventoryService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_duplicate_sale_lines_are_consolidated(self):
        pid = self.products.create_product("Duplicado venta", "Unidad", 2.0, 0)
        self.entries.register_entry(pid, 10, "2099-12-31", self.admin.id)
        sale_id = self.sales.create_sale([
            {"product_id": pid, "quantity": 3},
            {"product_id": pid, "quantity": 4},
        ], self.admin.id)
        lot = self.inventory.lots_for_product(pid)[0]
        self.assertEqual(lot["cantidad_actual"], 3)
        sale, details = self.sales.get_sale_details(sale_id)
        self.assertEqual(len(details), 1)
        self.assertEqual(details[0]["cantidad"], 7)
        self.assertEqual(sale["total"], 14.0)

    def test_duplicate_sale_lines_cannot_oversell(self):
        pid = self.products.create_product("Sin sobreventa", "Unidad", 1.0, 0)
        self.entries.register_entry(pid, 10, "2099-12-31", self.admin.id)
        with self.assertRaises(ValueError):
            self.sales.create_sale([
                {"product_id": pid, "quantity": 7},
                {"product_id": pid, "quantity": 7},
            ], self.admin.id)
        self.assertEqual(self.inventory.lots_for_product(pid)[0]["cantidad_actual"], 10)
        self.assertEqual(len(self.sales.list_sales()), 0)

    def test_restore_rejects_unrelated_sqlite_database(self):
        import sqlite3
        from app.services.backup_service import BackupService
        bogus = Path(self.tmp.name) / "bogus.db"
        con = sqlite3.connect(bogus)
        con.execute("CREATE TABLE unrelated(id INTEGER PRIMARY KEY)")
        con.commit()
        con.close()
        backup = BackupService(self.db)
        with self.assertRaises(ValueError):
            backup.restore_backup(bogus)
        self.assertIsNotNone(self.auth.login("admin", "Admin123!"))

    def test_unit_cannot_change_after_inventory_history(self):
        pid = self.products.create_product("Unidad estable", "Blíster", 1.0, 0)
        self.entries.register_entry(pid, 5, "2099-12-31", self.admin.id)
        with self.assertRaises(ValueError):
            self.products.update_product(pid, "Unidad estable", "Unidad", 1.0, 0, "")
        self.assertEqual(self.products.get_product(pid)["unidad_control"], "Blíster")

    def test_duplicate_names_have_friendly_validation(self):
        self.products.create_product("Mismo producto", "Unidad", 1.0, 0)
        with self.assertRaises(ValueError):
            self.products.create_product("mismo producto", "Unidad", 1.0, 0)
        self.auth.create_user("Persona Uno", "operador1", "Clave123", "OPERADOR")
        with self.assertRaises(ValueError):
            self.auth.create_user("Persona Dos", "OPERADOR1", "Clave123", "OPERADOR")

class ProductLifecycleRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "lifecycle.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.inventory = InventoryService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_product_with_stock_cannot_be_deactivated(self):
        pid = self.products.create_product("No ocultar stock", "Unidad", 1.0, 0)
        self.entries.register_entry(pid, 2, "2099-12-31", self.admin.id)
        with self.assertRaises(ValueError):
            self.products.set_active(pid, False)
        self.assertTrue(bool(self.products.get_product(pid)["activo"]))

    def test_zero_stock_product_can_be_deactivated(self):
        pid = self.products.create_product("Desactivable", "Unidad", 1.0, 0)
        self.products.set_active(pid, False)
        self.assertFalse(bool(self.products.get_product(pid)["activo"]))

class InventoryBusinessRuleTests(unittest.TestCase):
    """Regresiones de las reglas que no deben perderse al cambiar la interfaz."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "rules.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.sales = SaleService(self.db)
        self.inventory = InventoryService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_entry_can_replenish_active_product_with_zero_stock(self):
        pid = self.products.create_product("Sin existencias", "Blíster", 2.5, 5)
        self.assertEqual(self.products.get_product(pid)["stock_total"], 0)
        result = self.entries.register_entry(pid, 12, "2099-12-31", self.admin.id)
        self.assertEqual(result["quantity"], 12)
        self.assertEqual(self.products.get_product(pid)["stock_total"], 12)

    def test_entry_rejects_inactive_product(self):
        pid = self.products.create_product("Inactivo entrada", "Unidad", 1.0, 0)
        self.products.set_active(pid, False)
        with self.assertRaises(ValueError):
            self.entries.register_entry(pid, 5, "2099-12-31", self.admin.id)

    def test_sale_rejects_active_product_without_saleable_stock(self):
        pid = self.products.create_product("Sin stock venta", "Unidad", 1.0, 0)
        with self.assertRaises(ValueError):
            self.sales.create_sale([{"product_id": pid, "quantity": 1}], self.admin.id)

    def test_entry_and_sale_quantities_must_be_positive(self):
        pid = self.products.create_product("Cantidad válida", "Unidad", 1.0, 0)
        for qty in (0, -1):
            with self.assertRaises(ValueError):
                self.entries.register_entry(pid, qty, "2099-12-31", self.admin.id)
        self.entries.register_entry(pid, 5, "2099-12-31", self.admin.id)
        for qty in (0, -1):
            with self.assertRaises(ValueError):
                self.sales.create_sale([{"product_id": pid, "quantity": qty}], self.admin.id)

    def test_fefo_uses_earliest_expiration_first(self):
        pid = self.products.create_product("FEFO", "Blíster", 3.0, 0)
        later = self.entries.register_entry(pid, 10, "2099-12-31", self.admin.id)
        earlier = self.entries.register_entry(pid, 4, "2098-01-01", self.admin.id)
        self.sales.create_sale([{"product_id": pid, "quantity": 6}], self.admin.id)
        lots = {int(r["id"]): int(r["cantidad_actual"]) for r in self.inventory.lots_for_product(pid)}
        self.assertEqual(lots[earlier["lot_id"]], 0)
        self.assertEqual(lots[later["lot_id"]], 8)

    def test_expired_stock_is_physical_but_not_saleable(self):
        pid = self.products.create_product("Doble stock", "Unidad", 1.0, 0)
        self.entries.register_entry(pid, 7, "2020-01-01", self.admin.id)
        self.entries.register_entry(pid, 3, "2099-01-01", self.admin.id)
        p = self.products.get_product(pid)
        self.assertEqual(int(p["stock_total"]), 10)
        self.assertEqual(int(p["stock_vendible"]), 3)
        row = next(r for r in self.inventory.inventory_rows() if int(r["id"]) == pid)
        self.assertEqual(int(row["stock_vencido"]), 7)
        with self.assertRaises(ValueError):
            self.sales.create_sale([{"product_id": pid, "quantity": 4}], self.admin.id)


class ConfigurableExpiryPolicyTests(unittest.TestCase):
    def setUp(self):
        from datetime import date, timedelta
        self.date = date
        self.timedelta = timedelta
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "expiry_policy.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.sales = SaleService(self.db)
        self.settings = SettingsService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_default_policy_blocks_fifteen_days_or_less(self):
        self.assertEqual(self.settings.get_block_days(), 15)
        pid = self.products.create_product("Bloqueo 15", "Unidad", 1.0, 0)
        blocked_date = (self.date.today() + self.timedelta(days=15)).isoformat()
        allowed_date = (self.date.today() + self.timedelta(days=16)).isoformat()
        self.entries.register_entry(pid, 4, blocked_date, self.admin.id)
        self.entries.register_entry(pid, 3, allowed_date, self.admin.id)
        product = self.products.get_product(pid)
        self.assertEqual(int(product["stock_total"]), 7)
        self.assertEqual(int(product["stock_vendible"]), 3)
        with self.assertRaises(ValueError):
            self.sales.create_sale([{"product_id": pid, "quantity": 4}], self.admin.id)
        self.sales.create_sale([{"product_id": pid, "quantity": 3}], self.admin.id)

    def test_policy_change_is_applied_immediately(self):
        pid = self.products.create_product("Configurable", "Unidad", 2.0, 0)
        soon = (self.date.today() + self.timedelta(days=10)).isoformat()
        self.entries.register_entry(pid, 5, soon, self.admin.id)
        self.assertEqual(int(self.products.get_product(pid)["stock_vendible"]), 0)
        self.settings.set_block_days(5)
        self.assertEqual(self.settings.get_block_days(), 5)
        self.assertEqual(int(self.products.get_product(pid)["stock_vendible"]), 5)
        sale_id = self.sales.create_sale([{"product_id": pid, "quantity": 1}], self.admin.id)
        self.assertGreater(sale_id, 0)

    def test_policy_range_is_validated(self):
        with self.assertRaises(ValueError):
            self.settings.set_block_days(-1)
        with self.assertRaises(ValueError):
            self.settings.set_block_days(91)


class FinalStockSemanticsTests(unittest.TestCase):
    def setUp(self):
        from datetime import date, timedelta
        self.date = date
        self.timedelta = timedelta
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "stock_semantics.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)
        self.inventory = InventoryService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_total_saleable_blocked_and_expired_are_separated(self):
        pid = self.products.create_product("Semántica stock", "Unidad", 2.0, 1)
        blocked = (self.date.today() + self.timedelta(days=10)).isoformat()
        saleable = (self.date.today() + self.timedelta(days=45)).isoformat()
        expired = (self.date.today() - self.timedelta(days=2)).isoformat()
        self.entries.register_entry(pid, 4, blocked, self.admin.id)
        self.entries.register_entry(pid, 5, saleable, self.admin.id)
        self.entries.register_entry(pid, 3, expired, self.admin.id)
        p = self.products.get_product(pid)
        self.assertEqual(int(p["stock_total"]), 12)
        self.assertEqual(int(p["stock_vendible"]), 5)
        self.assertEqual(int(p["stock_bloqueado"]), 4)
        self.assertEqual(int(p["stock_vencido"]), 3)
        row = next(r for r in self.inventory.inventory_rows() if int(r["id"]) == pid)
        self.assertEqual(int(row["stock_total"]), 12)
        self.assertEqual(int(row["stock_vendible"]), 5)
        self.assertEqual(int(row["stock_bloqueado"]), 4)
        self.assertEqual(int(row["stock_vencido"]), 3)

    def test_entry_becomes_non_cancellable_after_sale(self):
        pid = self.products.create_product("Entrada protegida", "Unidad", 2.0, 1)
        created = self.entries.register_entry(pid, 5, "2099-12-31", self.admin.id)
        SaleService(self.db).create_sale([{"product_id": pid, "quantity": 1}], self.admin.id)
        row = next(r for r in self.entries.list_entries() if int(r["id"]) == created["entry_id"])
        self.assertFalse(bool(row["puede_anular"]))



class ConvenienceRegressionTests(unittest.TestCase):
    """Regresiones de pequeñas acciones que evitan repetir trabajo."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "convenience.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None
        self.products = ProductService(self.db)
        self.entries = EntryService(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_entry_descriptive_metadata_can_be_corrected_without_changing_stock(self):
        pid = self.products.create_product("Metadatos", "Unidad", 1.5, 0)
        created = self.entries.register_entry(
            pid,
            8,
            "2099-12-31",
            self.admin.id,
            manufacturer_lot="FAB-001",
            notes="Recepción inicial",
        )
        self.entries.update_entry_metadata(created["entry_id"], "FAB-001-C", "Observación corregida")
        row = next(r for r in self.entries.list_entries() if int(r["id"]) == created["entry_id"])
        self.assertEqual(row["lote_fabricante"], "FAB-001-C")
        self.assertEqual(row["observacion"], "Observación corregida")
        self.assertEqual(int(row["cantidad_inicial"]), 8)
        self.assertEqual(int(row["cantidad_actual"]), 8)

class FinalAutomationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "final.db")
        self.db.initialize()
        self.auth = AuthService(self.db)
        self.auth.ensure_default_admin()
        self.admin = self.auth.login("admin", "Admin123!")
        assert self.admin is not None

    def tearDown(self):
        self.tmp.cleanup()

    def test_theme_and_inactivity_settings_persist(self):
        settings = SettingsService(self.db)
        self.assertEqual(settings.get_theme(), "dark")
        self.assertEqual(settings.get_inactivity_minutes(), 0)
        settings.set_theme("light")
        settings.set_inactivity_minutes(30)
        self.assertEqual(SettingsService(self.db).get_theme(), "light")
        self.assertEqual(SettingsService(self.db).get_inactivity_minutes(), 30)
        with self.assertRaises(ValueError):
            settings.set_theme("otro")
        with self.assertRaises(ValueError):
            settings.set_inactivity_minutes(241)

    def test_daily_start_snapshot_and_last_close_are_separate(self):
        from app.services.backup_service import BackupService
        service = BackupService(self.db)
        start_first = service.ensure_daily_automatic_backup()
        start_second = service.ensure_daily_automatic_backup()
        self.assertEqual(start_first, start_second)
        self.assertTrue(start_first.exists())

        close_first = service.automatic_backup()
        close_second = service.automatic_backup()
        self.assertEqual(close_first, close_second)
        self.assertTrue(close_first.exists())
        self.assertNotEqual(start_first, close_first)

        history = service.history(20)
        start_rows = [r for r in history if r["tipo"] == "AUTOMATICO_INICIO"]
        close_rows = [r for r in history if r["tipo"] == "AUTOMATICO_CIERRE"]
        self.assertEqual(len(start_rows), 1)
        self.assertEqual(len(close_rows), 1)
        self.assertEqual(service.health_check(), [])

    def test_automatic_rotation_keeps_manual_backups(self):
        from app.services.backup_service import BackupService
        service = BackupService(self.db)
        for day in range(1, 10):
            (service.backup_dir / f"backup_inicio_202610{day:02d}.db").write_bytes(b"snapshot")
        manual = service.backup_dir / "respaldo_manual_importante.db"
        manual.write_bytes(b"manual")
        service.prune_automatic(7)
        self.assertEqual(len(list(service.backup_dir.glob("backup_inicio_????????.db"))), 7)
        self.assertTrue(manual.exists())

    def test_audit_log_records_sensitive_action(self):
        from app.services.audit_service import AuditService
        audit = AuditService(self.db)
        audit.log("PRUEBA", user_id=self.admin.id, entity="TEST", entity_id=1, detail="ok")
        rows = audit.recent(5)
        self.assertEqual(rows[0]["accion"], "PRUEBA")
        self.assertEqual(rows[0]["usuario"], "Administrador")

if __name__ == "__main__":
    unittest.main()
